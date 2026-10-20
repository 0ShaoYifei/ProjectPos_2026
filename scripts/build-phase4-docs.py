from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,re

root=Path('deliverables/presentation')
slides=json.loads((root/'presentation_content.json').read_text(encoding='utf-8'))['slides']
questions=json.loads((root/'defense_content.json').read_text(encoding='utf-8'))

def doc_base(title,zh=False):
 d=Document();s=d.sections[0]
 s.page_width=Inches(8.268);s.page_height=Inches(11.693)
 s.top_margin=s.bottom_margin=Inches(.8);s.left_margin=s.right_margin=Inches(.85)
 s.header_distance=s.footer_distance=Inches(.35)
 for name,size,bold in [('Normal',10.5,False),('Title',23,True),('Heading 1',15,True),('Heading 2',12,True),('Heading 3',11,True)]:
  style=d.styles[name];style.font.name='Arial';style.font.size=Pt(size);style.font.bold=bold;style.font.color.rgb=RGBColor(0,0,0)
  fonts=style.element.get_or_add_rPr().get_or_add_rFonts();fonts.set(qn('w:eastAsia'),'Microsoft YaHei')
  pf=style.paragraph_format;pf.line_spacing=1.24 if zh else 1.12;pf.space_after=Pt(6);pf.widow_control=True
  if name.startswith('Heading'):pf.keep_with_next=True;pf.space_before=Pt(12)
  if name=='Title':pf.space_after=Pt(14)
  snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');style.element.get_or_add_pPr().append(snap)
 # Remove decorative title rules inherited from the runtime's default template.
 for style in d.styles:
  for border in style.element.xpath('.//w:pBdr'):border.getparent().remove(border)
 d.core_properties.title=title;d.core_properties.author=''
 d.add_paragraph(title,'Title')
 f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
 fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
 for r in f.runs:r.font.name='Arial';r.font.size=Pt(9)
 return d

def para(d,text='',style=None):
 p=d.add_paragraph(style=style)
 for i,part in enumerate(re.split(r'(`[^`]+`)',text)):
  run=p.add_run(part[1:-1] if part.startswith('`') else part)
  if part.startswith('`'):run.font.name='Consolas';run.font.size=Pt(9)
 return p

def label(d,name,value):
 p=para(d);p.add_run(name+' ').bold=True
 p.add_run(value);return p

def markdown(source,output,toc=False):
 lines=source.read_text(encoding='utf-8').splitlines();d=doc_base(lines[0][2:],True)
 started=False
 for line in lines[1:]:
  if not line.strip():continue
  if line.startswith('## ') and toc and not started:
   heading=d.add_paragraph('目录');heading.runs[0].bold=True;heading.runs[0].font.size=Pt(15)
   p=d.add_paragraph();f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'TOC \\o "1-1" \\h \\z');p._p.append(f)
   d.add_page_break();started=True
  if line.startswith('### '):para(d,line[4:],'Heading 2')
  elif line.startswith('## '):para(d,line[3:],'Heading 1')
  elif line.startswith('- '):para(d,line[2:],'List Bullet')
  else:para(d,line)
 if toc:
  update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');d.settings.element.append(update)
 d.save(root/output)
 return d

d=doc_base('POS Presentation Video Script')
para(d,'ProjectPos_2026 is an individual Software Engineering Practice project. This script follows the fifteen-slide English presentation and includes the actual application demonstration.')
words=sum(len(s['narration'].split())for s in slides)
label(d,'Timing',f'{words} spoken words. Plan approximately 7 minutes 30 seconds, including screen actions. At 120–130 words per minute, allow about one minute of operation pauses. Rehearse and adjust to any teacher-specified limit.')
label(d,'Scope','The narration paragraphs are the words to speak. Key points, action directions and transition cues are rehearsal aids. Do not read every instruction aloud.')
label(d,'Before recording','Confirm MySQL and one POS instance are running. Open the slides, the homepage at http://127.0.0.1:8080/ and the genuine GitHub repository. The homepage redirects to Products. A brief real startup clip may precede the timed explanation.')
label(d,'Recording data','Use REC-20261020-RICE and REC-20261020-MILK, initially ten units each at CNY 6.50 and 3.80. These products were created separately without resetting earlier records. Rehearsals change stock; use a fresh unique suffix for another full take.')
label(d,'Demonstration sequence','Products 10/10; add rice 2 and milk 1; show total 16.80; reject cash 16.79; accept 20.00; show change 3.20 and the new receipt; verify stock 8/9; reject return quantity 3; accept rice return 1 for 6.50; show refund ledger and stock 9/9; show GitHub commits and retained branches.')
label(d,'Video URL','[Pending upload]. No video recording or upload is claimed by these materials.')
elapsed=0
for i,s in enumerate(slides):
 if i%2==0:d.add_page_break()
 start=elapsed;elapsed+=s['seconds'];stamp=lambda t:f'{t//60}:{t%60:02d}'
 para(d,f"Slide {s['number']} {s['title']}",'Heading 1')
 label(d,'Planned window',f'{stamp(start)}–{stamp(elapsed)} ({s["seconds"]} seconds). Adjust after rehearsal.')
 label(d,'Key points','; '.join(s['keyPoints'])+'.')
 label(d,'Narration',s['narration'])
 label(d,'Screen actions',s['actions'])
 label(d,'Transition cue',s['transition'])
d.save(root/'Video_Script_EN.docx')

markdown(root/'Project_Explanation_ZH.md','Project_Explanation_ZH.docx',toc=True)
markdown(root/'Recording_Guide_ZH.md','Recording_Guide_ZH.docx')

d=doc_base('项目答辩问题与双语回答',True)
para(d,'下面二十个问题围绕现有代码、设计与真实证据。先用英语问题练习口头回答，再用中文解析检查理解。参考路径是定位线索，实际演示时应打开对应代码或日志。')
for i,q in enumerate(questions):
 if i and i%2==0:d.add_page_break()
 para(d,f'问题 {i+1}','Heading 1')
 p=para(d,q['question']);p.paragraph_format.keep_with_next=True
 for r in p.runs:r.bold=True
 label(d,'English answer',q['answer'])
 label(d,'中文解析',q['explanation'])
 p=label(d,'代码与证据',q['evidence'])
 for r in p.runs:r.font.size=Pt(9)
d.save(root/'Defense_Questions_Bilingual.docx')

assert len(slides)==15 and len(questions)==20
for file in ['Video_Script_EN.docx','Project_Explanation_ZH.docx','Recording_Guide_ZH.docx','Defense_Questions_Bilingual.docx']:
 check=Document(root/file);print(file,'paragraphs',len(check.paragraphs),'characters',sum(len(p.text)for p in check.paragraphs))
