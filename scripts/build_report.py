from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
from zipfile import ZipFile,ZIP_DEFLATED
import json,csv,hashlib
root=Path('deliverables/report');assets=root/'report_assets';source=Path('references/Tech Report Cover Page (3).docx')
identity=json.loads(Path('.local/student-info.json').read_text());doc=Document(source);md=[];figure_no=0;table_no=0
source_hash=hashlib.sha256(source.read_bytes()).hexdigest()
with ZipFile(source) as z:source_parts={n:z.read(n) for n in z.namelist()}
replacements={8:'Name: '+identity['name'],9:'NUIST ID: '+identity['nuist_id'],10:'SETU ID: '+identity['setu_id'],11:'Email: '+identity['email'],12:'GitHub URL: github.com/0ShaoYifei/ProjectPos_2026',13:'Demo Video URL: [Pending upload]',15:'10/ 20 /2026'}
for index,text in replacements.items():
 nodes=doc.paragraphs[index]._p.xpath('.//w:t');assert nodes;nodes[0].text=text
 for n in nodes[1:]:n.text=''
for para in list(doc.paragraphs)[17:]:
 if not para.text.strip():para._p.getparent().remove(para._p)
for name,size,bold in [('Report Body',10.5,False),('Report Chapter',16,True),('Report Section',12,True),('Report Caption',9,False),('Report Code',8,False),('Report Table',9,False)]:
 style=doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH) if name not in doc.styles else doc.styles[name]
 style.font.name='Consolas' if name=='Report Code' else 'Arial';style.font.size=Pt(size);style.font.bold=bold;style.font.color.rgb=RGBColor(0,0,0)
 style.paragraph_format.line_spacing=1.12;style.paragraph_format.space_after=Pt(6);style.paragraph_format.widow_control=True
 snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');style.element.get_or_add_pPr().append(snap)
 if name in ['Report Chapter','Report Section']:
  style.paragraph_format.keep_with_next=True;style.paragraph_format.space_before=Pt(12)
  outline=OxmlElement('w:outlineLvl');outline.set(qn('w:val'),'0' if name=='Report Chapter' else '1');style.element.get_or_add_pPr().append(outline)
 if name=='Report Caption':style.paragraph_format.space_before=Pt(3);style.paragraph_format.space_after=Pt(12)
 if name=='Report Code':style.paragraph_format.line_spacing=1.0;style.paragraph_format.space_after=Pt(2)
 if name=='Report Table':style.paragraph_format.line_spacing=1.05;style.paragraph_format.space_after=Pt(3)
for tocname in ['TOC 1','TOC 2']:
 ts=doc.styles[tocname] if tocname in doc.styles else doc.styles.add_style(tocname,WD_STYLE_TYPE.PARAGRAPH)
 ts.font.name='Arial';ts.font.size=Pt(9.5);ts.font.color.rgb=RGBColor(0,0,0);ts.paragraph_format.line_spacing=1.0;ts.paragraph_format.space_before=Pt(0);ts.paragraph_format.space_after=Pt(3)
 sn=OxmlElement('w:snapToGrid');sn.set(qn('w:val'),'0');ts.element.get_or_add_pPr().append(sn)
toc=doc.add_paragraph();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'TOC \\t "Report Chapter,1,Report Section,2" \\h \\z');toc._p.append(field)
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');doc.settings.element.append(update)
section=doc.add_section(WD_SECTION_START.NEW_PAGE);section.footer.is_linked_to_previous=False
section.page_width=Inches(8.2680556);section.page_height=Inches(11.6930556);section.left_margin=section.right_margin=Inches(1.25);section.top_margin=section.bottom_margin=Inches(1)
page_start=OxmlElement('w:pgNumType');page_start.set(qn('w:start'),'1');section._sectPr.append(page_start)
footer=section.footer.paragraphs[0];footer.style='Report Caption';footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
def h1(text):
 para=doc.add_paragraph(text,'Report Chapter');para.paragraph_format.page_break_before=True;md.append('# '+text+'\n')
def h2(text):doc.add_paragraph(text,'Report Section');md.append('## '+text+'\n')
def p(text):
 para=doc.add_paragraph(text,'Report Body');para.alignment=WD_ALIGN_PARAGRAPH.LEFT;md.append(text+'\n');return para
def bullets(values):
 for value in values:
  para=p(value);para.paragraph_format.left_indent=Inches(.12)
  md[-1]='- '+value+'\n'
def table(title,headers,rows,widths=None):
 global table_no
 table_no+=1;cap=doc.add_paragraph(f'Table {table_no}  {title}','Report Caption');cap.paragraph_format.keep_with_next=True
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 if widths is None:widths=[5.768/len(headers)]*len(headers)
 for col,width in zip(t.columns,widths):col.width=Inches(width)
 for cell,text,width in zip(t.rows[0].cells,headers,widths):cell.width=Inches(width);cell.text=text
 for values in rows:
  for cell,text,width in zip(t.add_row().cells,values,widths):cell.width=Inches(width);cell.text=str(text)
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
 t._tbl.tblPr.append(borders);margins=OxmlElement('w:tblCellMar')
 for edge in ['top','left','bottom','right']:
  e=OxmlElement('w:'+edge);e.set(qn('w:w'),'90');e.set(qn('w:type'),'dxa');margins.append(e)
 t._tbl.tblPr.append(margins)
 for i,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
  if i==0:trpr.append(OxmlElement('w:tblHeader'))
  for j,cell in enumerate(row.cells):
   cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'24464B' if i==0 else ('F3F7F7' if i%2==0 else 'FFFFFF'));cell._tc.get_or_add_tcPr().append(shading)
   for para in cell.paragraphs:
    para.style='Report Table';para.paragraph_format.keep_with_next=False
    if j==0 and len(cell.text)<14:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in para.runs:r.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(0,0,0);r.font.bold=i==0
 doc.add_paragraph('',style='Report Body').paragraph_format.space_after=Pt(2)
 md.append(f'Table {table_no} {title}\n\n| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*len(headers))+' |\n'+'\n'.join('| '+' | '.join(map(str,row))+' |' for row in rows)+'\n')
def fig(filename,caption,max_width=5.7,max_height=7.1):
 global figure_no
 figure_no+=1;path=assets/filename
 with Image.open(path) as im:w,h=im.size
 width=min(max_width,max_height*w/h)
 para=doc.add_paragraph(style='Report Body');para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.paragraph_format.line_spacing=1.0;para.paragraph_format.keep_with_next=True;para.paragraph_format.space_after=Pt(1)
 para.add_run().add_picture(str(path),width=Inches(width))
 c=doc.add_paragraph(f'Figure {figure_no}  {caption}','Report Caption');c.alignment=WD_ALIGN_PARAGRAPH.LEFT
 md.append(f'![Figure {figure_no} {caption}](report_assets/{filename})\n')
def code(text):
 for line in text.splitlines():doc.add_paragraph(line,'Report Code')
 md.append('```text\n'+text+'\n```\n')
h1('Chapter 1 Introduction')
p('ProjectPos_2026 is a small supermarket Point-of-Sale application developed for the individual Software Engineering Practice course project. It implements Process Sale and Handle Returns, with supporting product and inventory management. The application records cash sales, calculates exact totals and change, retains receipts, and records refunds while maintaining stock quantities.')
h2('1.1 Background and objectives')
p('A checkout must keep the quantities sold, the cash received and the stock on hand consistent. Manual calculations can produce incorrect change, while unlinked returns can allow the same purchase to be refunded more than once. The project addresses these problems through validated transactions and a receipt-based return process. A cashier uses English forms and receives a clear explanation when an operation cannot be completed.')
p('The principal academic objective is to connect use-case analysis, domain modeling, layered object-oriented design, implementation and verification. Functional objectives are to manage products, maintain inventory, calculate running totals, accept cash, calculate change, generate receipts, record transactions and process valid returns. Supporting objectives are to preserve actual evidence and use meaningful Git feature branches.')
h2('1.2 Scope and boundaries')
p('The scope is one local supermarket installation with a trusted cashier. Prices and payments use CNY. Product stock is an aggregate count of interchangeable units. Each return request handles one purchased product line and can be repeated for additional lines. Card-payment integration, separate tax or discount rules, supplier purchasing, user accounts, external accounting and cash-drawer hardware are outside the implementation. Collecting and refunding physical cash and assessing whether goods are resalable remain cashier responsibilities.')
h2('1.3 Technology and initial inspection')
p('The existing project was inspected before modification. It contained the Spring Boot entry point, Maven configuration, database properties and one context-load test. No business entities, repositories, services, controllers or templates were present, and the directory had no Git repository. The starter was extended in place. The supplied Spring Boot 4.1.1 parent and Java 17 compilation target were retained. Actual verification used Oracle JDK 21, Maven 3.8.5 through the wrapper, Thymeleaf, Spring Web MVC, Jakarta Validation, Spring Data JPA, Hibernate and MySQL 8.0.44. H2 is used only by the default tests.')

h1('Chapter 2 Part 1 Agile Development Process')
h2('2.1 Individual Scrum adaptation')
p('The project uses an individual adaptation of Scrum: the student maintains an ordered backlog, selects a manageable increment, implements and tests it, reviews the result and adjusts later work. No team members, stand-up meetings or stakeholder approvals are invented. The observed sequence was product and inventory management, cash sale processing, return handling and integration evidence, followed by quizzes and report preparation. Each code increment received actual verification and a preserved feature branch.')
p('A backlog item is verified when its behavior is implemented and supported by an executed test, direct database verification or an inspected document, as appropriate. Writing a file alone is insufficient. The progress chart counts verified items. The items are not assumed to require equal effort and are not represented as story points or hours.')
h2('2.2 Product backlog and user stories')
backlog=[
('P01','Products','As a cashier, I want to create and edit products so prices and barcodes are accurate.','Must','Verified'),
('P02','Inventory','As a cashier, I want to maintain stock and deactivate unavailable products.','Must','Verified'),
('P03','Sales','As a cashier, I want to add and remove cart items before payment.','Must','Verified'),
('P04','Sales','As a cashier, I want an exact running total so I can state the amount due.','Must','Verified'),
('P05','Payments','As a cashier, I want to record sufficient cash to complete the purchase.','Must','Verified'),
('P06','Payments','As a cashier, I want the correct change displayed.','Must','Verified'),
('P07','Receipts','As a customer, I want a receipt preserving the original purchase details.','Must','Verified'),
('P08','Records','As a cashier, I want to review completed transactions.','Must','Verified'),
('P09','Returns','As a cashier, I want to find the receipt before accepting a return.','Must','Verified'),
('P10','Returns','As a customer, I want refunds calculated using the original price.','Must','Verified'),
('P11','Inventory','As a cashier, I want accepted returned units restored to stock.','Must','Verified'),
('P12','Validation','As a cashier, I want invalid transactions rejected with English messages.','Must','Verified'),
('P13','Quality','As a developer, I want repeatable normal, exceptional and concurrency tests.','Must','Verified'),
('P14','Database','As a developer, I want real MySQL integration and table evidence.','Must','Verified'),
('P15','Evidence','As a student, I want genuine browser and database evidence.','Must','Verified'),
('P16','GitHub','As a student, I want preserved Git branches and a confirmed remote repository.','Must','Local complete; remote pending'),
('P17','Quizzes','As a student, I want two complete source-mapped quizzes.','Must','Verified'),
('P18','Report','As a student, I want an editable report tied to actual evidence.','Must','This deliverable')]
with Path('docs/product-backlog.csv').open('w',newline='',encoding='utf-8') as f:
 writer=csv.writer(f);writer.writerow(['id','epic','user_story','priority','status']);writer.writerows(backlog)
table('Product backlog',['ID','Epic','User story','Priority','Result'],backlog,[.42,.65,2.52,.68,1.498])
p('P05 acceptance requires insufficient cash to leave stock and the sale ledger unchanged. P10 and P11 require an original-price refund, a stock increase and a refund record in one transaction. P12 covers invalid quantities, invalid prices, excessive returns and stale product edits. P16 remains open until a project remote is confirmed and all preserved branches are pushed.')
h2('2.3 Sprint planning and increment records')
p('The assignment suggests three-day or one-week sprints. The following three-day schedule is a planning baseline established on 20 October 2026 for reviewing and completing the coursework. It does not claim that three multi-day sprints have elapsed. Implementation increments were completed in the single recorded session on 20 October; later timeboxes remain proposed review and submission work.')
table('Proposed three day timeboxes',['Window','Goal','Status at report preparation'],[
('20–22 Oct 2026','Products, sales, returns and technical verification','Code and verification completed on 20 Oct; student review remains.'),
('23–25 Oct 2026','Review quizzes, UML and report','Documents prepared early; the scheduled review is proposed.'),
('26–28 Oct 2026','Presentation, video upload and GitHub submission','Proposed; completed external submissions are not claimed.')],[1.15,2.1,2.518])
table('Observed implementation increments',['Increment','Verified output','Evidence'],[
('Product','Product forms, inventory and stale edit checks','5 passing H2 tests; f716d77'),
('Sale','Cart, totals, cash, change and receipts','13 passing H2 tests; 5183641'),
('Return','Refunds, restocking and exception handling','21 passing H2 tests; a1fb0aa'),
('Integration','Real database and browser behavior','21 MySQL tests; 13 browser checks; SQL and screenshots'),
('Documents','Quizzes and current technical report','30 answers, source map, UML and rendered documents')],[.9,2.55,2.318])
p('Development reviews used executable feedback. The offline baseline exposed a missing Maven dependency, one MySQL command failed because PowerShell parsed an unquoted argument, and the first browser test used an unsuitable dropdown selector. Failed attempts remain beside corrected passing runs. These records support adjustment but are not evidence of unrecorded meetings or user acceptance.')
h2('2.4 Recorded burndown and limitations')
progress=list(csv.DictReader((assets/'burndown_data.csv').open(encoding='utf-8-sig')))
fig('10-recorded-burndown.png','Observed remaining work. Editable measurements are in burndown_data.csv.',max_height=3.3)
table('Actual progress checkpoints',['Time UTC plus 8','Verified items','Remaining items'],[(r['timestamp'][11:19],r['completed_work_items'],r['remaining_work_items'])for r in progress],[2.3,1.734,1.734])
p('The initial backlog contained 18 items. The chart covers only recorded timestamps from the baseline through quiz verification. Connecting lines do not imply that intermediate completion times were measured. At its cutoff, GitHub completion and the report remained open; this report was prepared afterward. There are no historical daily measurements, completed multi-day burndowns, velocity statistics or effort estimates to reconstruct. The editable CSV and dated log support future observations without invented history.')

h1('Chapter 3 Part 2 Requirements and Use Case Analysis')
h2('3.1 Functional requirements and business rules')
p('The cashier manages unique barcodes, names, positive unit prices, stock quantities and availability flags. A sale contains one to one hundred products; each requested quantity is between 1 and 10000. Prices have at most two decimal places. Stock lies between zero and one million. Cash must cover the total, and failed checkout records no sale. A receipt retains its original names and prices after catalog changes.')
p('Returns require a completed sale and a matching purchased line. Quantity must be positive and cannot exceed the purchased quantity minus all previous returns. A reason of one to two hundred characters is required. Refunds use the unit price captured in SaleItem. Sales reduce stock; accepted returns increase it without exceeding the inventory limit. Persisted request keys prevent another sale or refund for a repeated completed request.')
h2('3.2 Use case context')
fig('01-use-cases.png','Use case context. The cashier is the direct application actor.',max_height=5.2)
p('The customer participates in the business process by providing goods, cash and a receipt for returns. The cashier operates the software. Product management and transaction review support the required core use cases. The boundary includes the local database; no external card, accounting or inventory service is presented as implemented.')
h2('3.3 Fully dressed use case UC1 Process Sale')
p('Scope: ProjectPos_2026. Level: user goal. Primary actor: Cashier. Trigger: a customer presents goods. Stakeholders: the cashier needs correct totals and change; the customer needs accurate charges and proof of purchase; the store needs consistent stock and an auditable record.')
p('Preconditions: the application and database are available and the catalog contains the purchased products. The cashier has access to the trusted local installation. Authentication is not implemented. Minimal guarantee: failed checkout creates no sale and does not reduce stock. Success guarantee: one paid sale and its snapshots are saved, stock is reduced, cash and change are recorded, and a receipt is displayed.')
bullets(['1. The customer arrives at checkout with products.','2. The cashier opens New sale and selects a product and quantity.','3. The system checks the cart quantity and displays item details and the running total.','4. The cashier repeats entry, removing an item or clearing the cart if needed.','5. The cashier states the total, collects cash and enters the amount received.','6. The system locks products in a consistent order and revalidates availability, quantities, stock and cash.','7. The system records payment, persists the sale and snapshots, and deducts stock in one transaction.','8. The receipt and change are displayed. The cashier provides the receipt and change, and the customer leaves with the goods.'])
p('Extensions: unknown or inactive products are rejected. Insufficient stock or payment produces an English message and retains the cart for correction. Invalid numbers are rejected. Checkout uses current locked product information if the catalog changes. Database failures roll back the transaction. A stale browser checkout form is rejected, and a repeated service checkout key returns the existing sale. The saved receipt can be printed again without another sale.')
p('Special requirements: exact decimal money, historical line snapshots and English UI messages. Technology/data variations: web-form entry and browser printing or Save as PDF. Frequency: repeated per customer, with no measured throughput asserted. Open issues: authentication and dedicated receipt-printer integration.')
h2('3.4 Fully dressed use case UC2 Handle Returns')
p('Scope: ProjectPos_2026. Level: user goal. Primary actor: Cashier. Trigger: a customer presents purchased goods and the sale reference. Stakeholders: the customer needs a correct refund; the store needs to prevent excessive or duplicate refunds and maintain stock; the cashier needs the remaining returnable quantity.')
p('Preconditions: the application and database are available and the sale is completed. The cashier determines whether the goods can return to stock. Minimal guarantee: an invalid return changes neither stock nor refund records. Success guarantee: a return record is saved, the line’s returned quantity increases, stock is restored and the refund is displayed.')
bullets(['1. The cashier opens Transactions and returns and enters the sale number or chooses a receipt.','2. The system displays original items, prices, purchased quantities and quantities still returnable.','3. The cashier selects one purchased line and enters the return quantity and reason.','4. The system locks the line, checks receipt ownership and verifies the remaining quantity.','5. The product is locked and refund = original unit price x accepted quantity is calculated.','6. The system restores stock, updates returned quantity and saves the refund in one transaction.','7. The updated receipt and refund are displayed. The cashier provides the cash refund.','8. The process can be repeated for another purchased line.'])
p('Extensions: unknown receipts and items belonging to another receipt are rejected. Zero, negative, excessive and cumulatively excessive quantities, missing reasons, invalid tokens and inventory overflow are rejected. Concurrent attempts to return the same last unit are serialized by the line lock. Reusing a completed return key returns the saved refund without another stock increase. A database failure rolls back all return changes.')
p('Special requirements: use the original price even after renaming, repricing or deactivation. Technology/data variations: one product line per submission, with cash refunds only. Frequency: occasional relative to sales; no actual rate is asserted. Open issues: damaged-goods policy, return deadlines and managerial approvals.')
h2('3.5 System sequence diagrams')
p('The SSDs treat the application as a black box. Their operation names express cashier intent and are not claims that identically named Java methods exist. Internal collaborations appear in Chapter 6.')
fig('02-ssd-process-sale.png','Process Sale SSD with stock and payment alternatives.',max_height=5.5)
fig('03-ssd-handle-returns.png','Handle Returns SSD including invalid quantity rejection.',max_height=5.2)
h1('Chapter 4 Domain Model')
h2('4.1 Conceptual classes')
p('The model identifies four business concepts. Product describes a catalog item and its aggregate available quantity. Sale represents a completed cash purchase. SaleItem represents a product quantity and historical unit price within a sale. ReturnRecord represents an accepted return against a purchased line, with quantity, refund, reason and time. These concepts follow the use-case vocabulary and remain distinct from technical controller and repository classes.')
fig('04-domain-model.png','Conceptual domain model showing business attributes and multiplicities.',max_height=4.6)
h2('4.2 Associations and interpretation')
p('Each completed Sale contains one or more SaleItems, and each SaleItem belongs to exactly one Sale. A Product can be referenced by zero or many sale items; every sale item identifies one product. A SaleItem supports zero or many ReturnRecords, and each return record belongs to one purchased line. An additional arithmetic invariant requires total returned quantity not to exceed purchased quantity; multiplicity alone cannot express it.')
p('Cash payment is represented by Sale attributes rather than a separate Payment entity. Stock is a quantity on Product instead of a separate inventory subsystem or one object per physical unit. These choices keep the course application simple. The conceptual diagram excludes JPA identifiers, repositories and Java method signatures; those appear in the design diagrams.')

h1('Chapter 5 Logical Architecture')
h2('5.1 MVC and responsibility boundaries')
p('Spring MVC separates presentation from business use cases. Thymeleaf templates present data and submit forms. Controllers bind requests, manage page flow and the session cart, and delegate to services. Services define business operations and transaction boundaries. Repositories provide persistence queries and locks through Spring Data JPA. Entities hold persistent state and related domain behavior. BusinessException resides in a shared exception package, so entities do not depend on services.')
p('This is relaxed layering: controllers read entity data obtained through services for rendering, and services use repositories and entities. Lower layers do not reference web controllers or templates. The diagram reflects actual dependencies instead of imposing a strict rule that source code does not follow.')
fig('05-package-model.png','Logical packages and their actual dependencies.',max_height=6.5)
h2('5.2 Architectural rationale')
p('Refund and inventory rules can be tested without web-page interaction. Changing layout does not require changing stock deduction, and persistence queries are not duplicated in templates. Constructor injection makes dependencies explicit. The architecture deliberately avoids separately deployed services, message brokers and a custom persistence framework.')
p('SaleService and ReturnService coordinate business use cases; Spring MVC controllers manage presentation flow. This respects the course distinction between web MVC and GRASP controllers. Disabling open-in-view requires detailed sale queries to load the items needed before rendering a receipt. The design remains understandable without hiding business transactions behind unnecessary abstractions.')

h1('Chapter 6 Object Oriented Design')
h2('6.1 Classes and responsibilities')
table('Important design responsibilities',['Class','Responsibility and key behavior'],[
('Product','Validate catalog values; decreaseStock and restoreStock enforce stock rules; version detects stale edits.'),
('Sale','Own sale items, add subtotals, record cash/change and derive the refunded total.'),
('SaleItem','Retain original name, barcode and price; calculate subtotal and remaining returnable quantity.'),
('ReturnRecord','Retain accepted quantity, original-price refund, reason and timestamp.'),
('ShoppingCart','Store product IDs, quantities and checkout key in the session, without managed JPA entities.'),
('ProductService','Create products and serialize edits using version and barcode checks.'),
('SaleService','Preview a cart and revalidate/persist a cash checkout atomically.'),
('ReturnService','Validate receipt ownership and quantities; refund and restore stock atomically.'),
('Controllers','Bind forms, delegate business work and render or redirect English pages.'),
('Repositories','Find/persist entities and acquire pessimistic locks for contested quantities.')],[1.2,4.568])
fig('08-design-entities.png','Design classes for all four persistent entities. Selected fields and operations are shown.',max_height=6.9)
fig('09-design-services.png','Selected Process Sale controller, service and repository dependencies.',max_height=5.0)
h2('6.2 Interaction diagrams')
p('SaleController enters the transactional checkout service. Product rows are locked in sorted ID order, reducing deadlock risk for carts containing the same products in different orders. The service checks for an existing checkout key before another transaction is recorded. Sale creates its SaleItems, payment is recorded and stock is reduced before the aggregate is flushed.')
fig('06-interaction-checkout.png','Internal checkout collaboration within one transaction.',max_height=5.7)
p('Returns use a locked SaleItem to serialize competing requests. After obtaining the lock, ReturnService rechecks the request key, validates remaining quantity, locks the product, restores stock and saves a ReturnRecord. In the diagram, web = ReturnController, returns = ReturnService, items = SaleItemRepository, products = ProductRepository and records = ReturnRepository. Short object-role labels keep the interaction readable.')
fig('07-interaction-return.png','Internal return collaboration using actual controller, service and repository roles.',max_height=5.9)
h2('6.3 Responsibility assignment and design decisions')
p('Information Expert appears in SaleItem.getSubtotal, which knows both price and quantity, and Sale.getRefundedTotal, which knows its purchased lines. Creator appears in Sale.addItem, which creates the SaleItems it contains. Use-case coordination belongs to services, while web controllers delegate business work. These assignments support cohesion and a small gap between business concepts and code.')
p('No custom Factory or Adapter pattern is claimed: there are no interchangeable third-party tax or payment providers requiring them. Framework object creation does not establish a student-authored GoF factory. Tests were added and executed with increments, but evidence does not establish a historical test-first red/green/refactor sequence, so strict TDD practice is not asserted.')
p('Pessimistic locks protect current stock and returnable quantities inside transactions. Product’s JPA version and the submitted form version address a different problem: preventing an old editing page from overwriting a later stock change. Both mechanisms can be explained directly in a course defense.')

h1('Chapter 7 Implementation')
h2('7.1 Configuration and persistence')
p('Maven retains the supplied Spring Boot parent and adds H2 only in test scope. Runtime dependencies provide Web MVC, validation, Thymeleaf, Spring Data JPA and MySQL Connector/J. The application targets pos_2026_db on localhost:3306. DB_PASSWORD supplies the password; the original plaintext credential was removed before the first commit. The example environment file contains placeholders. Maven 3.8.5 was executed through the wrapper.')
p('Hibernate update mode creates missing tables. The verified schema contains products, sales, sale_items and return_records, all using InnoDB. Sale cascades persistence to its items. Products are deactivated rather than deleted, retaining references for sales and returns. SaleItem snapshots preserve the purchase even after catalog changes.')
h2('7.2 Product management and inventory')
p('ProductForm provides declarative constraints and English field messages. ProductService also invokes domain validation so rules do not rely on the browser. A database constraint protects barcode uniqueness. Editing locks the row and compares the form’s version to the current version before changing name, price, stock or active status. An outdated editing form must be reloaded after another transaction changes the product.')
fig('11-products-detail.png','Actual product and inventory page detail. Original screenshot 01-products.png is retained.',max_height=3.0)
h2('7.3 Cash sale and receipt generation')
p('ShoppingCart combines repeated quantities and checks available stock. SaleService.preview obtains product information for the running total. Checkout rechecks an empty cart, invalid quantities, inactive/missing products, insufficient stock, invalid decimal cash and insufficient payment. BigDecimal avoids binary floating-point errors, including the 0.10 multiplied by three scenario verified in tests.')
p('Total is the sum of unit price multiplied by quantity, and change is cash paid minus total. Cash and total are capped at CNY 999999999.99. Successful checkout saves the sale and historical lines and deducts stock together, then redirects to a saved receipt and clears the cart. A stale browser key is rejected; a unique persisted key also prevents duplicate service checkout.')
fig('12-receipt-detail.png','Actual POS-1 receipt: total 16.80, cash 20.00 and change 3.20. Unaltered detail crop.',max_height=4.8)
h2('7.4 Return processing and error handling')
p('A cashier finds a receipt, selects its purchased line and enters quantity and reason. ReturnService checks ownership and cumulative remaining quantity. Refund is original sale price multiplied by accepted quantity, even after repricing or deactivation. It restores stock, increments returnedQuantity and saves ReturnRecord in one transaction. The page displays the refund and separate return history.')
fig('13-return-detail.png','Actual return form and refund ledger after one rice unit was returned for CNY 6.50.',max_height=6.3)
p('Controller advice renders business and binding errors in English; product errors remain beside their fields. A fixed English locale avoids system-language messages. Each changing POST carries a session CSRF token checked by an interceptor. Database errors produce a friendly conflict page. This is a trusted localhost demonstration; account authentication and permission management remain future work.')
h1('Chapter 8 Testing and Results')
h2('8.1 Actual environment and test separation')
p('Verification ran on Windows 11 with Oracle JDK 21 and the Maven wrapper. Default tests use H2 in MySQL compatibility mode. The separate mysql-test profile runs the same service tests against real pos_2026_test_db, avoiding automated writes to the application’s pos_2026_db. Browser verification exercised the running application on 127.0.0.1:8080 against the application database. Clearly labeled DEMO products and their genuine sale/refund records remain for inspection.')
table('Executed build and integration results',['Execution','Observed result','Evidence in docs/evidence'],[
('Baseline offline clean test','Failed on uncached Surefire dependency; compilation completed','2026-10-20-baseline-maven-test.log'),
('Product increment','5 tests, zero failures/errors/skips','2026-10-20-products-final-test.log'),
('Sale clean test','13 tests, zero failures/errors/skips','2026-10-20-sales-test.log'),
('Complete clean test H2','21 tests, zero failures/errors/skips','2026-10-20-returns-test.log'),
('MySQL test profile','21 tests, zero failures/errors/skips','2026-10-20-mysql-integration-retry.log'),
('Clean package','BUILD SUCCESS; 21 passing tests; runnable JAR created','2026-10-20-clean-package.log'),
('spring-boot:run','Started on port 8080 and identified MySQL 8.0.44','2026-10-20-application-run.log'),
('Browser verification','13 checks passed; no JavaScript page errors','live-http-results.json')],[1.25,2.05,2.468])
p('The 21 tests comprise one context test, four product tests, eight sale tests, six return tests and two concurrency tests. Running them on two databases does not produce 42 distinct scenarios. Build success is not proof of every production behavior; the functional results below define the scope of verification.')
h2('8.2 Functional and exceptional scenarios')
table('Verified business scenarios',['Scenario','Expected result','Observed result'],[
('Multi-product sale','2 rice at 6.50 plus 1 milk at 3.80 totals 16.80','POS-1 saved with total 16.80.'),
('Change','20.00 cash minus 16.80 total','Receipt shows 3.20 change.'),
('Insufficient stock','Reject 10000 units without cart acceptance','English error; screenshot 02.'),
('Insufficient payment','Reject 16.79 for the 16.80 cart','English error; cart retained; screenshot 04.'),
('Failed checkout','No new sale or inventory reduction','Service assertions passed.'),
('Excessive return','Reject 3 units against 2 purchased','English error; screenshot 06.'),
('Partial refund','Refund 1 rice unit at original 6.50','RET-1 records refund 6.50.'),
('Inventory','Rice 30 - 2 + 1 = 29; milk 20 - 1 = 19','SQL and product page confirm 29 and 19.'),
('Cumulative over-return','Reject quantities beyond the remaining purchase','Automated assertions passed.'),
('Historical price','Use 2.35 after catalog price becomes 9.00','Two-unit refund verified as 4.70.'),
('Duplicate request','No second refund or restock with the same key','Service and live checks passed.'),
('Concurrency','Only one request consumes or refunds the last unit','Two-thread tests passed on H2 and MySQL.'),
('Forms and CSRF','Reject invalid values and missing token','English errors; missing token returned HTTP 403.'),
('Receipt lookup','Find saved receipt; reject malformed/unknown IDs','Lookup passed; invalid requests returned HTTP 400.')],[1.2,2.25,2.318])
h2('8.3 Real database evidence')
p('Direct mysql queries verified MySQL 8.0.44, pos_2026_db and four InnoDB tables. The saved sale has total 16.80, cash_paid 20.00 and change_due 3.20. Its rice line has quantity 2 and returned_quantity 1; milk has quantity 1 and returned_quantity 0. RET-1 references the rice line, with quantity 1 and refund 6.50. Product stock is rice 29, milk 19 and apples 25. These are saved demo records, not mocked query output.')
p('Results are retained in 2026-10-20-database-final-verification.log. Baseline SQL establishes that the schema already existed but had no tables. The initial sandboxed connection was denied by the environment; the authorized retry connected successfully. An unexecuted or failed check is not reported as a pass.')
h2('8.4 Corrections and verification limits')
p('The offline baseline lacked a Surefire dependency. The wrapper was aligned with available Maven 3.8.5, dependencies were downloaded and later tests/package passed. PowerShell split the first unquoted MySQL profile argument; the quoted retry passed. A dropdown locator caused the initial browser timeout; selecting the observed select element corrected the test. Escaped line labels were corrected before the final successful PlantUML render.')
p('Windows held an active server log open during a branch change. The POS server was stopped briefly, its active log was retained and the verified branch was merged. The packaged document renderer could not find LibreOffice in the Windows runtime. Installed Microsoft Word successfully exported the documents, and their PDF pages were rasterized for visual inspection.')
p('Passed: the executed automated suites, packaging, SQL and listed browser checks. Blocked: project-specific GitHub push and video URL confirmation. Not executed: load testing, long-duration reliability testing, penetration testing, hardware printing/cash-drawer tests and external payment integrations. No coverage percentage or unmeasured performance result is claimed.')

h1('Chapter 9 Git and GitHub Version Control')
h2('9.1 Initial state and individual authorship')
p('Initial status, branch, remote and log commands reported that the directory was not a Git repository. No earlier local history was available there. Git was initialized on 20 October 2026 after credentials were removed from tracked configuration, and the sanitized starter became the baseline. Commits use the student’s existing configured Git identity. This is individual development; no collaborators are invented.')
p('The GitHub connector authenticated as 0ShaoYifei but returned no accessible owned ProjectPos_2026 repository. There is no origin remote locally. The prior report’s GitHub link belongs to the robot project and is not used as the POS repository. Project repository URL: https://github.com/0ShaoYifei/ProjectPos_2026. No remote push or screenshot of a completed POS GitHub repository is claimed.')
h2('9.2 Preserved branches and merges')
table('Preserved implementation history',['Branch','Commit evidence','State'],[
('main','9f3d709, d129a15, a570cfb','Contains verified product, sale and return merges.'),
('feature/product-management','f716d77; 9f587a0','Preserved and merged with a merge commit.'),
('feature/process-sale','5183641','Preserved and merged with a merge commit.'),
('feature/handle-returns','a1fb0aa','Preserved and merged with a merge commit.'),
('feature/documentation','Current documents and evidence','Preserved; final delivery log records its merge.')],[1.65,1.65,2.468])
p('Commit messages describe concrete increments. Product work adds catalog and stock validation; the sale commit adds atomic cash checkout and receipts; the return commit adds audited refunds and real database/browser evidence. Merge commits preserve these boundaries. Branches have not been deleted, and there has been no force push or history rewriting.')
h2('9.3 Genuine command evidence and remote handoff')
p('Command journals contain actual commands and exit codes. The first journal includes a failed checkout caused by its own tracked log changing. The active journal was subsequently kept in ignored storage and snapshotted into evidence. A later active-server file-lock failure is also recorded. The branch and graph listings are direct Git output, not reconstructed history.')
code('git init -b main\ngit switch -c feature/product-management\ngit commit -m "feat: manage products and inventory with validation and stale edit protection"\ngit merge --no-ff feature/product-management ...\ngit switch -c feature/process-sale\ngit merge --no-ff feature/process-sale ...\ngit switch -c feature/handle-returns\ngit merge --no-ff feature/handle-returns ...')
p('These abbreviated commands summarize executed actions. Full commands and results appear in git-commands.log and git-session-record.log. Once the correct POS repository is confirmed, its remote history must be inspected before attaching origin and pushing main and all preserved feature branches. Any existing history must be integrated without overwriting it. Course reference files, local credentials and temporary working files are excluded from Git.')

h1('Chapter 10 Conclusion')
p('The application supports both required use cases with catalog management, exact cash calculations, receipts, transaction records, validated refunds and consistent inventory updates. Business concepts map to focused Java entities, transactional services and explicit persistence queries. Actual tests demonstrate normal, invalid and contested operations on H2 and real MySQL; browser and SQL evidence confirms the running system’s behavior.')
p('Incremental work exposed failures at manageable stages. Product rules were verified before checkout, and checkout before refunds and concurrent access. Dated records and retained failures support a defensible account of this work. Progress measurements cover one session, so they do not establish multi-day Scrum velocity or earlier sprint performance.')
p('Limitations include trusted local deployment, no authentication, no separate tax/discount or card subsystem, and a simple policy returning accepted goods to stock. Future work can add roles, damaged-goods handling, deadlines, pagination, migration-controlled schema changes and hardware receipt support. Immediate coursework actions are to confirm the POS remote, push the preserved history, prepare and review the presentation/video, upload the genuine demonstration and replace the pending project/video links before submission.')

h1('Appendix A Source and Evidence Index')
p('Authoritative specification: references/Software Engineering Practice Assignment.docx. Cover basis: references/Tech Report Cover Page (3).docx. Conceptual reference: Craig Larman, Applying UML and Patterns, third edition, especially chapters 6, 9, 10, 13, 15, 17, 19, 21 and 26. The previous robot report informed organization and, at the student’s explicit direction, identity fields. Its project results, dates, Git commands and progress are not POS evidence.')
table('Submission and evidence files',['Location','Contents'],[
('README.md','Setup, configuration, supported behavior and test instructions.'),
('src/main and src/test','Actual implementation and 21 automated tests.'),
('docs/evidence','Successful and failed logs, SQL, browser results, screenshots and dated development records.'),
('docs/product-backlog.csv','Editable backlog of 18 tracked items.'),
('report_assets/burndown_data.csv','Actual measurements frozen at report preparation.'),
('report_assets/*.puml, *.png, *.svg','Editable UML sources and rendered diagrams.'),
('screenshot-provenance.json','Original screenshot paths, hashes and unaltered crop coordinates.'),
('deliverables/quizzes','Both quizzes in Word/PDF and question-to-source mapping.'),
('report_evidence_checklist.md','Verification checks and outstanding external submission items.')],[2.05,3.718])
h2('A.1 Actual implementation excerpts')
p('The following excerpts are taken from the implementation; surrounding statements are omitted for clarity.')
code('''// SaleService.checkout
if (cash.compareTo(sale.getTotal()) < 0)
    throw new BusinessException("Insufficient payment. Cash must cover the sale total.");
sale.recordPayment(cash);
locked.forEach((id, product) -> product.decreaseStock(quantities.get(id)));
return sales.saveAndFlush(sale);

// ReturnService.process
product.restoreStock(quantity);
item.recordReturn(quantity);
return returns.saveAndFlush(new ReturnRecord(requestKey, item, quantity, reason));''')
h2('A.2 Demonstration and external submission fields')
p('POS repository URL: https://github.com/0ShaoYifei/ProjectPos_2026. Demonstration video URL: [Pending upload]. Replace these with genuine project-specific links. The earlier project video is not a substitute. A demonstration should combine slides and the working system: a cash sale and change, insufficient stock/payment, a valid refund, excessive-return rejection, database records and Git history. The assignment requires a presentation and Bilibili upload; this report does not claim those external submissions have already occurred.')
output=root/'POS_Technical_Report.docx';doc.core_properties.title='Analysis and Design of a PoS System';doc.core_properties.author=identity['name'];doc.save(output)
allowed={'word/document.xml','word/styles.xml','word/settings.xml','word/_rels/document.xml.rels','[Content_Types].xml','docProps/core.xml','docProps/app.xml'}
with ZipFile(output) as z:parts={n:z.read(n) for n in z.namelist()}
for name,data in source_parts.items():
 if name not in allowed:parts[name]=data
with ZipFile(output,'w',ZIP_DEFLATED) as z:
 for name,data in parts.items():z.writestr(name,data)
assert hashlib.sha256(source.read_bytes()).hexdigest()==source_hash
with ZipFile(output) as z:
 for name,data in source_parts.items():
  if name not in allowed:assert z.read(name)==data,name
(root/'POS_Technical_Report.md').write_text('\n'.join(md),encoding='utf-8')
inventory={'reference_sha256':source_hash,'preserved_parts':[{'path':n,'bytes':len(data),'sha256':hashlib.sha256(data).hexdigest()}for n,data in source_parts.items()if n not in allowed],'editable_parts':sorted(allowed),'figures':figure_no,'tables':table_no,'original_cover_logo_preserved':True}
Path('docs/evidence/report-template-preservation.json').write_text(json.dumps(inventory,indent=2),encoding='utf-8')
print('Report saved',output,'figures',figure_no,'tables',table_no,'paragraphs',len(doc.paragraphs))
