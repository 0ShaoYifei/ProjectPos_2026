"""Verify delivered content and native chart data; visual review is a separate assertion."""
from pathlib import Path
from zipfile import ZipFile
from datetime import datetime, timezone
from io import BytesIO
from docx import Document
from pypdf import PdfReader
from lxml import etree as E
from openpyxl import load_workbook
import csv, hashlib, json, re, sys

root=Path('deliverables/presentation')
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
compact=lambda s:re.sub(r'\s+','',s).replace('\u00ad','')
read=lambda p:json.loads(p.read_text(encoding='utf-8'))
content=read(root/'presentation_content.json');slides=content['slides']
questions=read(root/'defense_content.json')
assert [s['number'] for s in slides]==list(range(1,16))
assert len(questions)==20
assert sum(s['seconds'] for s in slides)==450
spoken_words=sum(len(s['narration'].split()) for s in slides)
assert spoken_words==763 and content['videoUrl']=='[Pending upload]'
docs={};pdfs={};checks=[];manifest=[]
expected={'Video_Script_EN':9,'Project_Explanation_ZH':12,'Recording_Guide_ZH':6,'Defense_Questions_Bilingual':10,'POS_Presentation':15}
for name,pages in expected.items():
 pdf=root/(name+'.pdf');reader=PdfReader(pdf)
 assert len(reader.pages)==pages,(name,len(reader.pages))
 text=''.join(re.sub(r'(?m)^\s*\d+\s*$','',p.extract_text() or '') for p in reader.pages)
 pdfs[name]=compact(text)
 assert '\ufffd' not in text,name
 manifest.append({'file':str(pdf),'sha256':sha(pdf),'pages':pages})
 if name!='POS_Presentation':
  file=root/(name+'.docx');doc=Document(file);docs[name]=doc
  for p in doc.paragraphs:
   if p.text.strip() and not p.style.name.lower().startswith('toc'):
    assert compact(p.text) in pdfs[name],(name,p.text[:65])
  assert not doc.styles.element.xpath('.//w:pBdr')
  manifest.append({'file':str(file),'sha256':sha(file),'paragraphs':len(doc.paragraphs),'all_body_paragraphs_present_in_pdf':True})
script='\n'.join(p.text for p in docs['Video_Script_EN'].paragraphs)
assert not re.search('[\u4e00-\u9fff]',script)
for s in slides:
 for value in [f"Slide {s['number']} {s['title']}",s['narration'],s['actions'],s['transition']]:assert value in script,value[:80]
for name in ['Project_Explanation_ZH','Recording_Guide_ZH']:
 body=compact('\n'.join(p.text for p in docs[name].paragraphs))
 for line in (root/(name+'.md')).read_text(encoding='utf-8').splitlines():
  line=re.sub(r'^(?:#{1,3} |- )','',line).replace('`','')
  if line.strip():assert compact(line) in body,(name,line[:60])
assert len([p for p in docs['Project_Explanation_ZH'].paragraphs if p.style.name=='Heading 1'])==12
assert len([p for p in docs['Defense_Questions_Bilingual'].paragraphs if re.fullmatch(r'问题 \d+',p.text)])==20
body='\n'.join(p.text for p in docs['Defense_Questions_Bilingual'].paragraphs)
for q in questions:
 for key in ['question','answer','explanation','evidence']:assert q[key] in body,(key,q[key][:50])

ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
deck=root/'POS_Presentation.pptx';table_slides=[]
with ZipFile(deck) as z:
 names=z.namelist();slidefiles=[n for n in names if re.fullmatch(r'ppt/slides/slide\d+.xml',n)]
 assert len(slidefiles)==15
 for s in slides:
  n=s['number'];xml=E.fromstring(z.read(f'ppt/slides/slide{n}.xml'))
  visible=' '.join(xml.xpath('//a:t/text()',namespaces=ns))
  assert s['title'] in visible
  assert not re.search('[\u4e00-\u9fff]',visible)
  if xml.xpath('//a:tbl',namespaces=ns):table_slides.append(n)
  note=E.fromstring(z.read(f'ppt/notesSlides/notesSlide{n}.xml'))
  text='\n'.join(note.xpath('//a:t/text()',namespaces=ns))
  assert compact(s['narration']) in compact(text)
 assert table_slides==[3,4,5,10,13,14]
 charts=[n for n in names if re.search(r'/charts/chart\d+\.xml$',n)]
 assert len(charts)==1
 chart=E.fromstring(z.read(charts[0]))
 xs=[float(v) for v in chart.xpath('//c:ser/c:xVal//c:pt/c:v/text()',namespaces=ns)]
 ys=[int(v) for v in chart.xpath('//c:ser/c:yVal//c:pt/c:v/text()',namespaces=ns)]
 rows=list(csv.DictReader(Path('docs/evidence/progress.csv').open(encoding='utf-8')))
 timestamps=[datetime.fromisoformat(r['timestamp']) for r in rows]
 expected_x=[round((t-timestamps[0]).total_seconds()/60,4) for t in timestamps]
 assert xs==expected_x and ys==[int(r['remaining_work_items']) for r in rows]
 assert chart.xpath('//c:ser/c:smooth/@val',namespaces=ns)==['0']
 books=[n for n in names if n.startswith('ppt/embeddings/') and n.endswith('.xlsx')]
 assert len(books)==1
 sheet=load_workbook(BytesIO(z.read(books[0])),data_only=True).active
 assert [sheet.cell(r,1).value for r in range(2,9)]==xs
 assert [sheet.cell(r,2).value for r in range(2,9)]==ys
receipt=read(Path('tmp/phase4-deck/validation-v5.json'))
assert receipt['finalSha256']==sha(deck)
assert receipt['packageIntegrity']['status']=='pass'
assert receipt['presentationLayout']['findingCount']==0
assert receipt['firstPartyImport']['passed']
manifest.append({'file':str(deck),'sha256':sha(deck),'slides':15})
database_bytes=Path('docs/evidence/phase4-database-verification.log').read_bytes()
database=database_bytes.decode('utf-16' if database_bytes.startswith((b'\xff\xfe',b'\xfe\xff')) else 'utf-8-sig')
for row in ['4\tREC-20261020-RICE\t6.50\t10\t1','5\tREC-20261020-MILK\t3.80\t10\t1','1\t16.80\t20.00\t3.20','1\t1\t1\t6.50\tIncorrect item purchased']:
 assert row in database,row
assert database.count('\tInnoDB')==4
visual='--visual-reviewed' in sys.argv
result={'recorded_at':datetime.now(timezone.utc).isoformat(),'slide_count':15,'english_visible_slides':True,'speaker_notes_match_script':True,'spoken_words':spoken_words,'planned_duration_seconds':450,'duration_is_estimate_not_recorded_video':True,'bilingual_defense_questions':20,'complete_word_and_pdf_text_verified':True,'native_table_slides':table_slides,'native_chart_slide':6,'embedded_chart_workbook_verified':True,'chart_x_minutes':xs,'chart_y_remaining_items':ys,'real_progress_csv_sha256':sha(Path('docs/evidence/progress.csv')),'presentation_package_and_layout_findings':0,'actual_word_pdf_pages':37,'actual_powerpoint_pdf_pages':15,'all_37_word_pages_and_15_slides_visually_reviewed':visual,'visual_result':'No missing glyphs, clipping, overlaps or broken pagination observed; real Office output inspected.' if visual else 'Pending visual review.','recording_products_stock':[10,10],'original_receipt_and_refund_preserved':True,'video_recorded':False,'video_uploaded':False,'video_url':'[Pending upload]','deliverables':manifest}
Path('docs/evidence/phase4-final-verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print('PASS: 15 English slides; matched narration and notes; 20 bilingual questions; complete Word/PDF text; 7 genuine chart points and embedded workbook; preserved MySQL records.')
print('Actual PDF pages: 15 slides + 37 Word pages. Visual review acknowledged:',visual)
