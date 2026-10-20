from pathlib import Path
from docx import Document
from docx.shared import Pt,RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,re,hashlib
root=Path('deliverables/quizzes');content=json.loads((root/'quiz_content.json').read_text());identity=json.loads(Path('.local/student-info.json').read_text())
verification=[]
for number in [1,2]:
 source=Path(f'references/Quiz{number}.docx');doc=Document(source)
 questions=[p.text for p in doc.paragraphs if re.match(r'^\d+\.',p.text.strip())]
 assert len(questions)==15
 normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10.5);normal.font.color.rgb=RGBColor(0,0,0)
 normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.1
 for p in doc.paragraphs:
  p.paragraph_format.line_spacing=1.1;p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(7);p.paragraph_format.keep_with_next=False;p.paragraph_format.page_break_before=False;p.paragraph_format.widow_control=True
 for role in ['Quiz Question','Quiz Answer']:
  if role not in doc.styles:doc.styles.add_style(role,WD_STYLE_TYPE.PARAGRAPH)
  style=doc.styles[role];style.base_style=normal;style.font.name='Arial';style.font.size=Pt(11 if role=='Quiz Question' else 10.5);style.font.color.rgb=RGBColor(0,0,0)
  style.font.bold=role=='Quiz Question';style.paragraph_format.keep_with_next=role=='Quiz Question'
  style.paragraph_format.space_before=Pt(11 if role=='Quiz Question' else 0);style.paragraph_format.space_after=Pt(5 if role=='Quiz Question' else 9)
  style.paragraph_format.widow_control=True
 title=doc.paragraphs[0];title.text=f'Software Engineering Practice Quiz{number}';title.style=doc.styles['Title']
 title.paragraph_format.space_after=Pt(14)
 for r in title.runs:r.font.name='Arial';r.font.size=Pt(17);r.font.bold=True;r.font.color.rgb=RGBColor(0,0,0)
 fields={'NuistID:':'NuistID: '+identity['nuist_id'],'SEUID:':'SEUID: '+identity['setu_id'],'Name:':'Name: '+identity['name']}
 for p in doc.paragraphs[1:4]:
  p.text=fields[p.text.strip()];p.paragraph_format.space_after=Pt(3)
 paragraphs=doc.paragraphs;answered=[]
 for i,p in enumerate(paragraphs):
  match=re.match(r'^(\d+)\.',p.text.strip())
  if not match:continue
  index=int(match.group(1));p.style=doc.styles['Quiz Question'];p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(10)
  # Keep the original question wording and runs, including the supplied typo.
  for r in p.runs:r.font.name='Arial';r.font.size=Pt(11);r.font.bold=True;r.font.color.rgb=RGBColor(0,0,0)
  target=paragraphs[i+1];assert not target.text.strip();target.text=content[f'quiz{number}'][index-1];target.style=doc.styles['Quiz Answer']
  for r in target.runs:r.font.name='Arial';r.font.size=Pt(10.5);r.font.color.rgb=RGBColor(0,0,0)
  answered.append(index)
 footer=doc.sections[0].footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
 for s in doc.sections:s.header_distance=Pt(28);s.footer_distance=Pt(28)
 doc.core_properties.author='';doc.core_properties.title=f'Software Engineering Practice Quiz{number} Answers';doc.core_properties.subject='Complete answers to fifteen supplied questions'
 output=root/f'Quiz{number}_Answers.docx';doc.save(output)
 reopened=Document(output);found=[p.text for p in reopened.paragraphs if re.match(r'^\d+\.',p.text.strip())]
 assert found==questions;assert answered==list(range(1,16))
 verification.append({'quiz':number,'questions':len(found),'answers':len(answered),'wording_preserved':True,'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest()})
Path('docs/evidence/quiz-structure-check.json').write_text(json.dumps(verification,indent=2),encoding='utf-8')
print(json.dumps(verification,indent=2))
